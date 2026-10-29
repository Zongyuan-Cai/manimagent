from __future__ import annotations

import io
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from manimagent.api import create_app
from manimagent.budget import Budget
from manimagent.config import ModelConfig, Settings
from manimagent.content_store import ContentError, ContentStore, new_id
from manimagent.markdown_view import render_markdown
from manimagent.providers import ChatProvider, ProviderError
from manimagent.schemas import Generation, Job, TextGenerationRequest, Usage
from manimagent.text_jobs import TextJobManager
from tests.support import wait_job


def settings_for(tmp_path, **kwargs):
    return Settings(
        runs_dir=tmp_path / "runs",
        models={
            "text": ModelConfig(
                base_url="https://model.example/team/v1", model="text-model", auth_mode="none"
            )
        },
        **kwargs,
    )


def controlled_provider(settings, calls, decision="pass", finish="stop", invalid_ref=False):
    def handler(request):
        body = json.loads(request.content)
        calls.append(body)
        if "response_format" in body:
            content = json.dumps({"decision": decision, "message": "检查结果"})
        else:
            payload = json.loads(body["messages"][1]["content"])
            ref = payload["allowed_source_refs"][0]
            if invalid_ref:
                ref["block_id"] = new_id("block")
            content = (
                "# 投影解答\n\n投影系数是 $3$，投影向量为 $(3,0)$。\n\n"
                "\\[\\operatorname{proj}_{b}(a)=(3,0)\\]\n\n"
                f"[[{ref['material_id']}:{ref['block_id']}]]\n"
            )
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": content}, "finish_reason": finish}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 50},
            },
        )

    return ChatProvider(settings, httpx.MockTransport(handler))


def submit(client, inputs=None, **source_kwargs):
    return client.post(
        "/api/v2/generations",
        json={
            "source": {
                "type": "inputs",
                "inputs": inputs or [{"type": "text", "text": "a=(3,2), b=(1,0)"}],
                "instruction": "求投影并解释",
                **source_kwargs,
            }
        },
    )


def get_generation_answer(client, job):
    generation = client.get(f"/api/v2/generations/{job['resource_ref']['id']}").json()
    answer = client.get(f"/api/v2/answers/{generation['answer_id']}").json()
    return generation, answer


def pdf_bytes(texts):
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in texts:
        page = writer.add_blank_page(width=600, height=800)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode("ascii") if text else b"")
        page[NameObject("/Contents")] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def test_complete_text_answer_revision_reuse_and_rendering(tmp_path):
    settings = settings_for(tmp_path)
    calls = []
    with TestClient(
        create_app(settings, text_provider=controlled_provider(settings, calls))
    ) as client:
        response = submit(client)
        assert response.status_code == 202
        job = wait_job(client, response.json()["job_id"])
        assert job["status"] == "completed", job
        generation, answer = get_generation_answer(client, job)
        assert generation["output_mode"] == "text" and set(generation["stages"]) == {
            "material",
            "answer",
        }
        assert len(calls) == 2
        assert (
            "response_format" not in calls[0]
            and "只返回符合" not in calls[0]["messages"][0]["content"]
        )
        assert "response_format" in calls[1]
        assert "candidate_markdown" in json.loads(calls[1]["messages"][1]["content"])
        assert job["usage"]["model_calls"] == 2
        markdown = client.get(f"/api/v2/answers/{answer['answer_id']}/markdown?download=true")
        assert (
            markdown.text
            == answer["markdown"]
            == "".join(step["content"] for step in answer["steps"])
        )
        assert "attachment" in markdown.headers["content-disposition"]
        html = client.get(f"/api/v2/answers/{answer['answer_id']}/html").text
        assert "<h1>投影解答</h1>" in html and r"\operatorname" in html
        revised = client.post(
            "/api/v2/answers/revisions",
            json={
                "base_answer_id": answer["answer_id"],
                "markdown": answer["markdown"] + "\n核对完成。",
            },
        )
        assert revised.status_code == 201
        assert revised.json()["derived_from_answer_id"] == answer["answer_id"]
        assert client.get(f"/api/v2/answers/{answer['answer_id']}").json() == answer
        reused = client.post(
            "/api/v2/generations",
            json={
                "source": {
                    "type": "materials",
                    "material_ids": answer["material_ids"],
                    "instruction": "换个解释",
                }
            },
        )
        reuse_job = wait_job(client, reused.json()["job_id"])
        assert reuse_job["status"] == "completed"
        assert get_generation_answer(client, reuse_job)[0]["material_ids"] == answer["material_ids"]
        assert len(calls) == 4


def test_upload_and_parse_work_without_any_models_or_docker(tmp_path, monkeypatch):
    settings = Settings(runs_dir=tmp_path)
    with TestClient(create_app(settings)) as client:
        files = [
            client.post(
                "/api/v2/files",
                files={"file": ("../same.md", "# 题设\n\na=(3,2)".encode(), "image/png")},
            ).json()
            for _ in range(2)
        ]
        assert files[0]["filename"] == "same.md" and files[0]["media_type"] == "text/markdown"
        assert (
            files[0]["file_id"] != files[1]["file_id"] and files[0]["sha256"] == files[1]["sha256"]
        )
        assert (
            client.get(f"/api/v2/files/{files[0]['file_id']}/content").content
            == "# 题设\n\na=(3,2)".encode()
        )
        response = client.post(
            "/api/v2/materials/parse",
            json={
                "inputs": [
                    {"type": "text", "text": "\n直接输入\n\n第二段"},
                    *[{"type": "file", "file_id": f["file_id"]} for f in files],
                ]
            },
        )
        job = wait_job(client, response.json()["job_id"])
        assert job["status"] == "completed", job
        material = client.get(f"/api/v2/materials/{job['result']['id']}").json()
        assert material["blocks"][0]["source"]["locator"]["line_start"] == 2
        assert [b["source"]["input_index"] for b in material["blocks"]] == [0, 0, 1, 1, 2, 2]
        assert material["blocks"][0]["source"]["file_id"] is None
        rejected = submit(client)
        assert (
            rejected.status_code == 503 and rejected.json()["error"]["code"] == "model_unavailable"
        )


@pytest.mark.parametrize(
    "filename, data, expected",
    [
        ("x.bin", b"abc", 415),
        ("x.txt", b"\xff", 415),
        ("x.pdf", b"not a pdf", 415),
        ("x.txt", b"\0", 415),
        ("x.md", b"x" * 20, 413),
    ],
)
def test_upload_limits_and_formats_leave_no_published_asset(tmp_path, filename, data, expected):
    settings = settings_for(tmp_path, max_upload_bytes=10)
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/v2/files", files={"file": (filename, data)})
        assert response.status_code == expected, response.text
        assert "error" in response.json()
        assert not list((settings.runs_dir / "private" / "files").glob("*"))


def test_chunked_multipart_limit_and_multiple_files(tmp_path):
    with TestClient(create_app(settings_for(tmp_path, max_upload_bytes=4))) as client:
        response = client.post(
            "/api/v2/files",
            content=iter([b"x" * 40000, b"x" * 40000]),
            headers={"Content-Type": "multipart/form-data; boundary=test"},
        )
        # A malformed body may be rejected before the byte limit is reached.
        assert response.status_code in {400, 413}
        response = client.post(
            "/api/v2/files", files=[("file", ("a.txt", b"a")), ("file", ("b.txt", b"b"))]
        )
        assert response.status_code == 400 and "error" in response.json()


def test_pdf_selection_coverage_and_blank_pages(tmp_path):
    with TestClient(create_app(settings_for(tmp_path))) as client:
        asset = client.post(
            "/api/v2/files", files={"file": ("notes.pdf", pdf_bytes(["FIRST", "SECOND", ""]))}
        ).json()
        response = client.post(
            "/api/v2/materials/parse",
            json={
                "inputs": [
                    {
                        "type": "file",
                        "file_id": asset["file_id"],
                        "selection": {"kind": "pages", "start": 2, "end": 2},
                    }
                ]
            },
        )
        job = wait_job(client, response.json()["job_id"])
        assert job["status"] == "completed", job
        material = client.get(f"/api/v2/materials/{job['result']['id']}").json()
        assert material["blocks"][0]["content"] == "SECOND"
        assert material["blocks"][0]["source"]["locator"] == {"page": 2}
        response = client.post(
            "/api/v2/materials/parse",
            json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
        )
        job = wait_job(client, response.json()["job_id"])
        assert job["status"] == "needs_review" and job["result"] is None
        diagnostic = client.get(f"/api/v2/files/{job['diagnostic_asset_ids'][-1]}/content").json()
        assert diagnostic["parse"]["coverage"][0]["missing_ranges"] == [[3, 3]]


def test_docx_order_table_and_unsupported_equation(tmp_path):
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "a", "3"
    document.add_paragraph("After")
    stream = io.BytesIO()
    document.save(stream)
    with TestClient(create_app(settings_for(tmp_path))) as client:
        asset = client.post(
            "/api/v2/files", files={"file": ("notes.docx", stream.getvalue())}
        ).json()
        response = client.post(
            "/api/v2/materials/parse",
            json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
        )
        job = wait_job(client, response.json()["job_id"])
        assert job["status"] == "completed", job
        blocks = client.get(f"/api/v2/materials/{job['result']['id']}").json()["blocks"]
        assert [b["kind"] for b in blocks] == ["text", "table", "text"]
        assert blocks[0]["content"] == "Before" and blocks[2]["content"] == "After"
        assert "| a | 3 |" in blocks[1]["content"] and blocks[1]["source"]["locator"] is None
        from docx.oxml import OxmlElement

        document.paragraphs[0]._p.append(OxmlElement("m:oMath"))
        stream = io.BytesIO()
        document.save(stream)
        asset = client.post(
            "/api/v2/files", files={"file": ("formula.docx", stream.getvalue())}
        ).json()
        job = wait_job(
            client,
            client.post(
                "/api/v2/materials/parse",
                json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
            ).json()["job_id"],
        )
        assert job["status"] == "needs_review" and job["result"] is None


@pytest.mark.parametrize(
    "mode, decision, finish, bad_ref, expected",
    [
        ("normal", "needs_input", "stop", False, "needs_input"),
        ("normal", "needs_review", "stop", False, "needs_review"),
        ("normal", "pass", "length", False, "needs_review"),
        ("normal", "pass", "stop", True, "needs_review"),
    ],
)
def test_incomplete_answers_never_publish(tmp_path, mode, decision, finish, bad_ref, expected):
    settings = settings_for(tmp_path)
    calls = []
    with TestClient(
        create_app(
            settings, text_provider=controlled_provider(settings, calls, decision, finish, bad_ref)
        )
    ) as client:
        job = wait_job(client, submit(client).json()["job_id"])
        assert job["status"] == expected, job
        assert job["result"] is None and job["diagnostic_asset_ids"]
        generation = client.get(f"/api/v2/generations/{job['resource_ref']['id']}").json()
        assert generation["answer_id"] is None
        assert not list((settings.runs_dir / "private" / "answers").glob("*.json"))


def test_aggregate_context_budget_fails_before_model_call(tmp_path):
    settings = settings_for(tmp_path, answer_context_tokens=1024)
    calls = []
    with TestClient(
        create_app(settings, text_provider=controlled_provider(settings, calls))
    ) as client:
        job = wait_job(client, submit(client).json()["job_id"])
        assert job["error"]["code"] == "budget_exceeded" and not calls


@pytest.mark.parametrize(
    "body",
    [
        {},
        {"source": {"type": "inputs", "inputs": [], "instruction": "q"}},
        {
            "source": {
                "type": "inputs",
                "inputs": [{"type": "text", "text": "\ud800"}],
                "instruction": "q",
            }
        },
        {
            "source": {
                "type": "inputs",
                "inputs": [{"type": "text", "text": "q"}],
                "instruction": " ",
            }
        },
        {"source": {"type": "answer", "answer_id": "../../config.yaml"}},
        {"output_mode": "video", "source": {"type": "answer", "answer_id": new_id("answer")}},
    ],
)
def test_v2_validation_does_not_create_jobs(tmp_path, body):
    settings = settings_for(tmp_path)
    with TestClient(create_app(settings)) as client:
        response = client.post(
            "/api/v2/generations",
            content=json.dumps(body).encode(),
            headers={"Content-Type": "application/json"},
        )
        assert response.status_code == 422 and "error" in response.json()
        assert not list((settings.runs_dir / "private" / "jobs").glob("*.json"))


def test_missing_references_and_native_mode_are_rejected_before_queue(tmp_path):
    with TestClient(create_app(settings_for(tmp_path))) as client:
        response = submit(client, inputs=[{"type": "file", "file_id": new_id("file")}])
        assert response.status_code == 404
        asset = client.post("/api/v2/files", files={"file": ("a.txt", b"a")}).json()
        response = client.post(
            "/api/v2/materials/parse",
            json={
                "input_mode": "native",
                "inputs": [{"type": "file", "file_id": asset["file_id"]}],
            },
        )
        assert response.status_code == 422
        # Native mode never changes how direct text is handled.
        job = wait_job(
            client,
            client.post(
                "/api/v2/materials/parse",
                json={"input_mode": "native", "inputs": [{"type": "text", "text": "abc"}]},
            ).json()["job_id"],
        )
        assert job["status"] == "completed"


def test_queue_capacity_snapshot_and_restart(tmp_path):
    settings = settings_for(tmp_path, max_waiting=0)
    manager = TextJobManager(settings)
    request = TextGenerationRequest(
        source={"type": "inputs", "inputs": [{"type": "text", "text": "题目"}], "instruction": "q"}
    )
    job = manager.submit("generation", request)
    with pytest.raises(ContentError) as exc:
        manager.submit("generation", request)
    assert exc.value.http_status == 429
    assert len(list((manager.store.root / "jobs").glob("*.json"))) == 1
    settings.models["text"].model = "changed"
    record = json.loads(
        (manager.store.root / "records" / f"{job.job_id}.json").read_text(encoding="utf-8")
    )
    assert record["settings"]["models"]["text"]["model"] == "text-model"
    restarted = TextJobManager(settings)
    restarted.start()
    try:
        recovered = restarted.store.load("jobs", job.job_id, Job)
        assert recovered.status == "interrupted"
        generation = restarted.store.load("generations", job.resource_ref.id, Generation)
        assert generation.stages["answer"].status == "interrupted"
    finally:
        restarted.stop()


def test_reuse_saved_answer_needs_no_model_configuration(tmp_path):
    settings = settings_for(tmp_path)
    calls = []
    with TestClient(
        create_app(settings, text_provider=controlled_provider(settings, calls))
    ) as client:
        _, answer = get_generation_answer(client, wait_job(client, submit(client).json()["job_id"]))
    settings.models = {}
    with TestClient(create_app(settings)) as client:
        job = wait_job(
            client,
            client.post(
                "/api/v2/generations",
                json={"source": {"type": "answer", "answer_id": answer["answer_id"]}},
            ).json()["job_id"],
        )
        assert job["status"] == "completed" and job["usage"]["model_calls"] == 0
        assert get_generation_answer(client, job)[1] == answer


def test_markdown_html_and_math_are_safe_and_lossless():
    html = render_markdown(
        r"<script>alert(1)</script> $x<y$ [bad](javascript:alert(1))" + "\n\n\\[a=\\frac{1}{2}\\]"
    )
    assert "<script>" not in html and 'href="javascript:' not in html
    assert r"\(x&lt;y\)" in html and r"\[a=\frac{1}{2}\]" in html
    assert "<img" not in render_markdown("![描述](https://remote.example/image.png)")


def test_custom_auth_full_url_and_text_finish_errors(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXT_TEST_KEY", "private-test-key")
    settings = settings_for(tmp_path)
    cfg = settings.models["text"]
    cfg.auth_mode, cfg.api_key_env, cfg.api_key_header = (
        "api_key_header",
        "TEXT_TEST_KEY",
        "Gateway-Key",
    )
    cfg.generate_url = "https://model.example/custom/invoke"

    def handler(request):
        assert (
            str(request.url) == cfg.generate_url
            and request.headers["Gateway-Key"] == "private-test-key"
        )
        assert "authorization" not in request.headers
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "partial"}, "finish_reason": "length"}]}
        )

    provider = ChatProvider(settings, httpx.MockTransport(handler))
    with pytest.raises(ProviderError) as exc:
        provider.complete_text("answer_writer", "answer", {}, Budget(settings, Usage()))
    assert exc.value.code == "output_truncated"


def test_independent_answer_endpoints_and_cli(tmp_path, monkeypatch, capsys):
    from manimagent.cli import main

    settings = settings_for(tmp_path)
    calls = []
    provider = controlled_provider(settings, calls)
    with TestClient(create_app(settings, text_provider=provider)) as client:
        parsed = client.post(
            "/api/v2/materials/parse",
            json={"inputs": [{"type": "text", "text": "a=(3,2), b=(1,0)"}]},
        ).json()
        material_id = wait_job(client, parsed["job_id"])["result"]["id"]
        job = wait_job(
            client,
            client.post(
                "/api/v2/answers", json={"material_ids": [material_id], "instruction": "求投影"}
            ).json()["job_id"],
        )
        assert job["status"] == "completed" and job["result"]["type"] == "answer"
        answer = client.get(f"/api/v2/answers/{job['result']['id']}").json()
        rejected = client.post(
            "/api/v2/answers/revisions",
            json={"base_answer_id": answer["answer_id"], "markdown": "丢弃全部来源"},
        )
        assert rejected.status_code == 422
    settings = settings_for(tmp_path / "cli")
    monkeypatch.setattr("manimagent.cli.load_settings", lambda _: settings)
    monkeypatch.setattr(
        "manimagent.cli.TextJobManager",
        lambda cfg: TextJobManager(cfg, controlled_provider(cfg, calls)),
    )
    file = tmp_path / "notes.txt"
    file.write_text("a=(3,2), b=(1,0)", encoding="utf-8")
    output = tmp_path / "answer.md"
    assert (
        main(["answer", "--question", "求投影", "--file", str(file), "--output", str(output)]) == 0
    )
    assert "投影解答" in output.read_text(encoding="utf-8")
    assert not capsys.readouterr().err


def test_parser_timeout_and_archive_limit_preserve_original(tmp_path):
    settings = settings_for(tmp_path, parse_timeout=0.001)
    with TestClient(create_app(settings)) as client:
        asset = client.post(
            "/api/v2/files", files={"file": ("notes.pdf", pdf_bytes(["text"]))}
        ).json()
        job = wait_job(
            client,
            client.post(
                "/api/v2/materials/parse",
                json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
            ).json()["job_id"],
        )
        assert job["status"] == "failed" and job["error"]["code"] == "budget_exceeded"
        assert client.get(f"/api/v2/files/{asset['file_id']}/content").status_code == 200
    document = Document()
    document.add_paragraph("body")
    stream = io.BytesIO()
    document.save(stream)
    settings = settings_for(tmp_path / "archive", max_archive_entries=1)
    with TestClient(create_app(settings)) as client:
        asset = client.post(
            "/api/v2/files", files={"file": ("notes.docx", stream.getvalue())}
        ).json()
        job = wait_job(
            client,
            client.post(
                "/api/v2/materials/parse",
                json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
            ).json()["job_id"],
        )
        assert job["status"] == "needs_input" and job["error"]["code"] == "budget_exceeded"


def test_modified_original_is_rejected_and_failures_release_queue(tmp_path):
    settings = settings_for(tmp_path, max_waiting=0)

    def handler(request):
        return httpx.Response(401, text="do not echo private-key")

    with TestClient(
        create_app(settings, text_provider=ChatProvider(settings, httpx.MockTransport(handler)))
    ) as client:
        asset = client.post("/api/v2/files", files={"file": ("notes.txt", b"text")}).json()
        store = ContentStore(settings)
        store.file_path(asset["file_id"]).write_bytes(b"altered")
        response = client.post(
            "/api/v2/materials/parse",
            json={"inputs": [{"type": "file", "file_id": asset["file_id"]}]},
        )
        assert (
            response.status_code == 422 and response.json()["error"]["code"] == "source_incomplete"
        )
        for _ in range(2):
            job = wait_job(client, submit(client).json()["job_id"])
            assert (
                job["error"]["code"] == "provider_auth_error" and job["usage"]["model_calls"] == 1
            )
            assert "private-key" not in json.dumps(job)


def test_provider_unknown_usage_and_malformed_messages(tmp_path):
    settings = settings_for(tmp_path)

    def handler(request):
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "complete"}, "finish_reason": "stop"}]}
        )

    budget = Budget(settings, Usage())
    assert (
        ChatProvider(settings, httpx.MockTransport(handler)).complete_text(
            "answer_writer", "answer", {}, budget
        )
        == "complete"
    )
    assert budget.usage.model_calls == 1 and not budget.usage.tokens_reported

    def bad_handler(request):
        return httpx.Response(200, json={"choices": [{"message": []}]})

    with pytest.raises(ProviderError):
        ChatProvider(settings, httpx.MockTransport(bad_handler)).complete_text(
            "answer_writer", "answer", {}, Budget(settings, Usage())
        )
